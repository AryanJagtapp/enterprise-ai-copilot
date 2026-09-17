"""
Document parsing + chunking.

Supports PDF, DOCX, TXT (per spec). Parsing failures are classified and
raised as AppError so the ingestion pipeline can record a clean
"failed" status instead of crashing the request.
"""
import re
from dataclasses import dataclass
from typing import List

from app.core.errors import AppError, ErrorClass

CHUNK_SIZE_CHARS = 1200
CHUNK_OVERLAP_CHARS = 150


@dataclass
class Chunk:
    index: int
    text: str


def parse_txt(raw_bytes: bytes) -> str:
    try:
        return raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return raw_bytes.decode("latin-1", errors="replace")


def parse_pdf(raw_bytes: bytes) -> str:
    try:
        from io import BytesIO

        from pypdf import PdfReader
    except ImportError as exc:
        raise AppError(ErrorClass.DEPENDENCY_UNAVAILABLE, "PDF parsing is unavailable on this server.", detail=str(exc)) from exc

    try:
        reader = PdfReader(BytesIO(raw_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages)
    except Exception as exc:  # pypdf raises various exceptions for corrupt PDFs
        raise AppError(
            ErrorClass.INVALID_INPUT,
            "This PDF could not be read — it may be corrupted, encrypted, or scanned without OCR text.",
            detail=str(exc),
        ) from exc


def parse_docx(raw_bytes: bytes) -> str:
    try:
        from io import BytesIO

        from docx import Document as DocxDocument
    except ImportError as exc:
        raise AppError(ErrorClass.DEPENDENCY_UNAVAILABLE, "DOCX parsing is unavailable on this server.", detail=str(exc)) from exc

    try:
        doc = DocxDocument(BytesIO(raw_bytes))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise AppError(
            ErrorClass.INVALID_INPUT,
            "This DOCX file could not be read — it may be corrupted or not a valid Word document.",
            detail=str(exc),
        ) from exc


PARSERS = {
    ".txt": parse_txt,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
}


def parse_document(filename: str, raw_bytes: bytes) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    parser = PARSERS.get(ext)
    if parser is None:
        raise AppError(ErrorClass.INVALID_INPUT, f"Unsupported file type '{ext}'. Allowed: {list(PARSERS)}")
    text = parser(raw_bytes)
    if not text or not text.strip():
        raise AppError(ErrorClass.INVALID_INPUT, "No extractable text was found in this document.")
    return text


def chunk_text(text: str, *, chunk_size: int = CHUNK_SIZE_CHARS, overlap: int = CHUNK_OVERLAP_CHARS) -> List[Chunk]:
    """
    Simple, dependency-free sliding-window chunker on paragraph boundaries
    where possible, falling back to a hard character window. Deterministic
    and fast — appropriate for the corpus sizes this project targets;
    swapping in a token-aware chunker later is a drop-in change since
    callers only depend on the Chunk dataclass shape.
    """
    normalized = re.sub(r"\n{3,}", "\n\n", text.strip())
    paragraphs = [p.strip() for p in normalized.split("\n\n") if p.strip()]

    chunks: List[str] = []
    current = ""
    for para in paragraphs:
        candidate = f"{current}\n\n{para}" if current else para
        if len(candidate) <= chunk_size:
            current = candidate
        else:
            if current:
                chunks.append(current)
            if len(para) <= chunk_size:
                current = para
            else:
                # A single paragraph longer than chunk_size: hard-split with overlap.
                start = 0
                while start < len(para):
                    end = start + chunk_size
                    chunks.append(para[start:end])
                    start = end - overlap
                current = ""
    if current:
        chunks.append(current)

    if not chunks:
        chunks = [normalized[i : i + chunk_size] for i in range(0, len(normalized), chunk_size - overlap)]

    return [Chunk(index=i, text=c) for i, c in enumerate(chunks)]
